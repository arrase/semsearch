"""Document and text extractor supporting txt, md, html, docx, odt, and rtf."""

import re
import xml.etree.ElementTree as ET
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.table import Table, _Cell

from semsearch.core.search_text import has_words
from semsearch.core.types import ExtractedSegment, Modality, SegmentKind


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.pieces: list[str] = []

    def handle_data(self, data: str) -> None:
        self.pieces.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in ("p", "div", "br", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr"):
            self.pieces.append("\n")


def strip_html(html_text: str) -> str:
    """Strip HTML tags and convert to plain text."""
    parser = _HTMLTextExtractor()
    parser.feed(html_text)
    text = "".join(parser.pieces)
    lines = [line.strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _strip_rtf(rtf_text: str) -> str:
    """Strip RTF formatting control sequences to extract plain text."""
    pattern = re.compile(r"\{\\(?:fonttbl|colortbl|stylesheet|info|pict)[^{}]*\}", re.DOTALL)
    while pattern.search(rtf_text):
        rtf_text = pattern.sub("", rtf_text)
    rtf_text = re.sub(r"\\(?:par|line)\b", "\n", rtf_text)
    rtf_text = re.sub(r"\\\'([0-9a-fA-F]{2})", lambda m: chr(int(m.group(1), 16)), rtf_text)
    rtf_text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", rtf_text)
    rtf_text = re.sub(r"[{}]", "", rtf_text)
    lines = [line.strip() for line in rtf_text.splitlines()]
    return "\n".join(line for line in lines if line)


def _read_odt(path: str) -> str:
    """Extract paragraph and heading text from OpenDocument Text (.odt) zip archive."""
    with zipfile.ZipFile(path) as zf:
        content_xml = zf.read("content.xml")
    tree = ET.fromstring(content_xml)
    paragraphs: list[str] = []
    for elem in tree.iter():
        if elem.tag.endswith(("}p", "}h")):
            text = "".join(elem.itertext()).strip()
            if text:
                paragraphs.append(text)
    return "\n\n".join(paragraphs)


def _extract_docx_table_rows(table: Table) -> list[str]:
    seen_cells: set[_Cell] = set()
    rows: list[str] = []
    for row in table.rows:
        row_parts: list[str] = []
        for cell in row.cells:
            if cell in seen_cells:
                continue
            seen_cells.add(cell)
            text = cell.text.strip()
            if text:
                row_parts.append(text)
        if row_parts:
            rows.append("\t".join(row_parts))
    return rows


def _read_docx(path: str) -> str:
    """Extract paragraph and table cell text from Word (.docx) document."""
    doc = Document(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        parts.extend(_extract_docx_table_rows(table))
    return "\n\n".join(parts)


def read_doc_text(path: str) -> str:
    """Read plain text from various document formats (.txt, .md, .html, .docx, .odt, .rtf)."""
    p = Path(path)
    suffix = p.suffix.lower()

    if suffix in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="replace")
    if suffix in (".html", ".htm"):
        raw_html = p.read_text(encoding="utf-8", errors="replace")
        return strip_html(raw_html)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".odt":
        return _read_odt(path)
    if suffix == ".rtf":
        raw_rtf = p.read_text(encoding="utf-8", errors="replace")
        return _strip_rtf(raw_rtf)

    return p.read_text(encoding="utf-8", errors="replace")


def _make_chunk_segment(text: str, loc: int) -> ExtractedSegment:
    return ExtractedSegment(
        kind=SegmentKind.CHUNK,
        modality=Modality.TEXT,
        loc=float(loc),
        excerpt=text[:200],
        payload=text,
    )


def _snap_chunk_end(clean_text: str, start: int, end: int, overlap: int) -> int:
    search_start = max(start, end - overlap)
    for i in range(end - 1, search_start - 1, -1):
        if clean_text[i].isspace():
            return i if i > start else end
    return end


def _advance_chunk_start(clean_text: str, start: int, end: int, overlap: int) -> int:
    n = len(clean_text)
    next_start = max(start + 1, end - overlap)
    while next_start < end and not clean_text[next_start].isspace():
        next_start += 1
    while next_start < n and clean_text[next_start].isspace():
        next_start += 1
    return next_start if next_start > start else end


def chunk_document(text: str, chunk_size: int = 1800, overlap: int = 200) -> list[ExtractedSegment]:
    """Slice document text into chunks snapping to word boundaries and filtering non-text."""
    clean_text = text.strip()
    n = len(clean_text)
    if n == 0:
        return []

    if n <= chunk_size:
        return [_make_chunk_segment(clean_text, 0)] if has_words(clean_text) else []

    chunks: list[ExtractedSegment] = []
    start = 0
    chunk_idx = 0

    while start < n:
        end = min(start + chunk_size, n)
        if end < n:
            end = _snap_chunk_end(clean_text, start, end, overlap)

        chunk_str = clean_text[start:end].strip()
        if chunk_str and has_words(chunk_str):
            chunks.append(_make_chunk_segment(chunk_str, chunk_idx))
            chunk_idx += 1

        if end >= n:
            break

        start = _advance_chunk_start(clean_text, start, end, overlap)

    return chunks


def extract_doc(path: str, chunk_size: int = 1800, overlap: int = 200) -> list[ExtractedSegment]:
    """Extract document file and return list of ExtractedSegment objects."""
    text = read_doc_text(path)
    return chunk_document(text, chunk_size=chunk_size, overlap=overlap)
