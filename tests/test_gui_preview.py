"""Unit and widget tests for multimodal preview components."""

from pathlib import Path
import zipfile

from docx import Document
from PIL import Image
import pymupdf
import pytest
from PySide6.QtGui import QTextFormat
from PySide6.QtWidgets import QApplication, QPlainTextEdit

from semsearch.core.types import FileKind, SearchHit
from semsearch.gui.preview import (
    AudioPreviewWidget,
    CodePreviewWidget,
    ImagePreviewWidget,
    PdfPreviewWidget,
    PreviewPane,
    VideoPreviewWidget,
    format_file_size,
    format_tilde_path,
)


def test_format_helpers():
    """Verify size and tilde formatting helpers."""
    assert format_file_size(500) == "500 B"
    assert format_file_size(2048) == "2.0 KB"
    assert format_file_size(5 * 1024 * 1024) == "5.0 MB"

    home_path = str(Path.home() / "Documents" / "test.txt")
    assert format_tilde_path(home_path).startswith("~/")


def test_image_preview_widget(tmp_path: Path, qapp: QApplication):
    """Verify ImagePreviewWidget scales image and updates metadata."""
    img_path = tmp_path / "sample.png"
    img = Image.new("RGB", (400, 300), color="blue")
    img.save(img_path)

    widget = ImagePreviewWidget()
    widget.load_image(str(img_path), "sample.png", size=img_path.stat().st_size, mtime=img_path.stat().st_mtime)

    assert "400 × 300 px" in widget.dimensions_label.text()
    assert widget.image_label.pixmap() is not None
    assert not widget.image_label.pixmap().isNull()
    assert widget.meta_card.name_label.text() == "sample.png"


def test_pdf_preview_widget(tmp_path: Path, qapp: QApplication):
    """Verify PdfPreviewWidget renders page and highlights keywords with navigation."""
    pdf_path = tmp_path / "sample.pdf"
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "Semantic search in Linux desktop engine")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Second page content here")
    doc.save(str(pdf_path))
    doc.close()

    widget = PdfPreviewWidget()
    widget.load_pdf(str(pdf_path), "sample.pdf", page_number=1, query="search desktop")

    assert "Page 1 of 2" in widget.page_info.text()
    assert widget.page_label.pixmap() is not None
    assert not widget.page_label.pixmap().isNull()
    assert widget.next_btn.isEnabled()
    assert not widget.prev_btn.isEnabled()

    widget._next_page()
    assert "Page 2 of 2" in widget.page_info.text()
    assert not widget.next_btn.isEnabled()
    assert widget.prev_btn.isEnabled()


def test_audio_preview_widget(qapp: QApplication):
    """Verify AudioPreviewWidget waveform visualization and scrub updates."""
    widget = AudioPreviewWidget()
    widget.load_audio("/tmp/track.mp3", "track.mp3", loc=45.0, size=1024 * 1024)

    assert widget._duration >= 45.0
    assert widget.scrub_slider.value() > 0
    assert widget.waveform._progress > 0.0

    widget._on_scrub(500)
    assert widget.waveform._progress == pytest.approx(0.5, 0.01)


def test_video_preview_widget(qapp: QApplication):
    """Verify VideoPreviewWidget poster thumbnail and scrub."""
    widget = VideoPreviewWidget()
    widget.load_video("/tmp/clip.mp4", "clip.mp4", loc=30.0, size=2 * 1024 * 1024)

    assert widget.poster_label.pixmap() is not None
    assert not widget.poster_label.pixmap().isNull()
    assert widget.scrub_slider.value() > 0


def test_code_preview_widget(tmp_path: Path, qapp: QApplication):
    """Verify CodePreviewWidget loads code and configures line gutter and highlight."""
    code_path = tmp_path / "sample.py"
    code_path.write_text("def hello():\n    return 42\n\nprint(hello())\n", encoding="utf-8")

    widget = CodePreviewWidget()
    widget.load_code(str(code_path), "sample.py", line=2)

    assert "def hello():" in widget.editor.toPlainText()
    assert len(widget.editor.extraSelections()) == 1
    assert widget.editor.gutter_width() > 0


def test_preview_pane_stacked_routing(tmp_path: Path, qapp: QApplication):
    """Verify PreviewPane routes multimodal items correctly across stacked indices."""
    pane = PreviewPane()
    assert pane.currentIndex() == 0

    # Image hit
    img_path = tmp_path / "test.jpg"
    Image.new("RGB", (100, 100)).save(img_path)
    hit_img = SearchHit(path=str(img_path), kind=FileKind.IMAGE, score=1.0)
    pane.load_result(hit_img)
    assert pane.currentIndex() == 1

    # PDF hit
    pdf_path = tmp_path / "test.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 50), "PDF content")
    doc.save(str(pdf_path))
    doc.close()
    hit_pdf = SearchHit(path=str(pdf_path), kind=FileKind.PDF, score=1.0, loc=1.0)
    pane.load_result(hit_pdf, query="content")
    assert pane.currentIndex() == 2

    # Audio hit
    hit_audio = SearchHit(path="/tmp/music.mp3", kind=FileKind.AUDIO, score=1.0, loc=10.0)
    pane.load_result(hit_audio)
    assert pane.currentIndex() == 3

    # Video hit
    hit_video = SearchHit(path="/tmp/movie.mp4", kind=FileKind.VIDEO, score=1.0, loc=20.0)
    pane.load_result(hit_video)
    assert pane.currentIndex() == 4

    # Code hit
    py_path = tmp_path / "test.py"
    py_path.write_text("x = 10", encoding="utf-8")
    hit_code = SearchHit(path=str(py_path), kind=FileKind.CODE, score=1.0, focus=1.0)
    pane.load_result(hit_code)
    assert pane.currentIndex() == 5

    # Clear
    pane.clear()
    assert pane.currentIndex() == 0


def test_code_preview_widget_document_formats(tmp_path: Path, qapp: QApplication):
    """Verify CodePreviewWidget extracts text from binary documents (.docx, .odt) rather than raw bytes."""
    # 1. Non-existent file
    widget = CodePreviewWidget()
    widget.load_code(str(tmp_path / "nonexistent.docx"), "nonexistent.docx")
    assert widget.editor.toPlainText() == ""

    # 2. DOCX format
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Docx Header", level=1)
    doc.add_paragraph("Docx preview paragraph text.")
    doc.save(str(docx_path))

    widget.load_code(str(docx_path), "sample.docx")
    content = widget.editor.toPlainText()
    assert "Docx Header" in content
    assert "Docx preview paragraph text." in content
    assert "[Content_Types].xml" not in content

    # 3. ODT format
    odt_path = tmp_path / "sample.odt"
    content_xml = b"""<?xml version="1.0" encoding="UTF-8"?>
    <office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
                             xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
      <office:body>
        <office:text>
          <text:p>ODT extracted body text.</text:p>
        </office:text>
      </office:body>
    </office:document-content>"""
    with zipfile.ZipFile(odt_path, "w") as zf:
        zf.writestr("content.xml", content_xml)

    widget.load_code(str(odt_path), "sample.odt")
    odt_content = widget.editor.toPlainText()
    assert "ODT extracted body text." in odt_content
    assert "office:document-content" not in odt_content

    # 4. Uppercase extension (.DOCX)
    upper_docx = tmp_path / "UPPER.DOCX"
    doc_upper = Document()
    doc_upper.add_paragraph("Uppercase DOCX content.")
    doc_upper.save(str(upper_docx))
    widget.load_code(str(upper_docx), "UPPER.DOCX")
    assert "Uppercase DOCX content." in widget.editor.toPlainText()

    # 5. RTF format
    rtf_path = tmp_path / "sample.rtf"
    rtf_path.write_text(r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Courier;}}\f0\fs20 RTF extracted body text.\par}", encoding="utf-8")
    widget.load_code(str(rtf_path), "sample.rtf")
    assert "RTF extracted body text." in widget.editor.toPlainText()

    # 6. HTML format
    html_path = tmp_path / "sample.html"
    html_path.write_text("<html><body><h1>HTML Heading</h1><p>HTML paragraph text.</p></body></html>", encoding="utf-8")
    widget.load_code(str(html_path), "sample.html")
    html_content = widget.editor.toPlainText()
    assert "HTML Heading" in html_content
    assert "HTML paragraph text." in html_content
    assert "<h1>" not in html_content


def test_preview_pane_document_routing(tmp_path: Path, qapp: QApplication):
    """Verify PreviewPane routes document files to code/text preview with extracted text and only term highlights."""
    docx_path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("Annual report content.")
    doc.save(str(docx_path))

    pane = PreviewPane()

    hit = SearchHit(path=str(docx_path), kind=FileKind.DOC, score=1.0)
    pane.load_result(hit, query="Annual")
    assert pane.currentIndex() == 5
    assert "Annual report content." in pane.code_view.editor.toPlainText()
    selections = pane.code_view.editor.extraSelections()
    assert len(selections) == 1
    assert not selections[0].format.boolProperty(QTextFormat.Property.FullWidthSelection)
    assert pane.code_view.editor.lineWrapMode() == QPlainTextEdit.LineWrapMode.WidgetWidth


def test_preview_pane_doc_and_code_highlight_behavior(tmp_path: Path, qapp: QApplication):
    """Verify FileKind.DOC (.md) omits full-width yellow line highlight while FileKind.CODE with line > 0 retains it."""
    md_path = tmp_path / "Bienvenido.md"
    md_path.write_text("Esta es su nueva *bóveda*.\n\nSegunda línea con contenido sobre un pozo profundo.\n", encoding="utf-8")

    pane = PreviewPane()

    # Semantic match on .md where query term 'agujero' is not present literally
    hit_md = SearchHit(
        path=str(md_path),
        kind=FileKind.DOC,
        score=0.8,
        excerpt="Segunda línea con contenido sobre un pozo profundo.",
    )
    pane.load_result(hit_md, query="agujero")
    assert pane.currentIndex() == 5
    assert pane.code_view.editor.extraSelections() == []
    assert pane.code_view.editor.lineWrapMode() == QPlainTextEdit.LineWrapMode.WidgetWidth
    assert pane.code_view.editor.textCursor().blockNumber() == 2

    # Code match with focus line > 0 retains full-width yellow line highlight and NoWrap
    py_path = tmp_path / "engine.py"
    py_path.write_text("import os\n\ndef compute_score():\n    return 1.0\n", encoding="utf-8")
    hit_code = SearchHit(path=str(py_path), kind=FileKind.CODE, score=0.9, focus=3.0)
    pane.load_result(hit_code, query="compute")
    code_selections = pane.code_view.editor.extraSelections()
    assert len(code_selections) == 2
    assert code_selections[0].format.boolProperty(QTextFormat.Property.FullWidthSelection)
    assert not code_selections[1].format.boolProperty(QTextFormat.Property.FullWidthSelection)
    assert pane.code_view.editor.lineWrapMode() == QPlainTextEdit.LineWrapMode.NoWrap


def test_metadata_card_actions(tmp_path: Path, qapp: QApplication, monkeypatch):
    """Verify MetadataCard Open and Show in Files action buttons and action_triggered signals."""
    opened = []
    revealed = []
    card_triggers = []
    pane_triggers = []
    monkeypatch.setattr("semsearch.gui.preview.open_file", lambda p: opened.append(p))
    monkeypatch.setattr("semsearch.gui.preview.reveal_in_file_manager", lambda p: revealed.append(p))

    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("Hello world", encoding="utf-8")

    pane = PreviewPane()
    pane.code_view.meta_card.action_triggered.connect(lambda: card_triggers.append(True))
    pane.action_triggered.connect(lambda: pane_triggers.append(True))
    hit = SearchHit(path=str(txt_path), kind=FileKind.DOC, score=1.0)
    pane.load_result(hit)

    pane.code_view.meta_card.open_btn.click()
    pane.code_view.meta_card.reveal_btn.click()

    assert opened == [str(txt_path)]
    assert revealed == [str(txt_path)]
    assert len(card_triggers) == 2
    assert len(pane_triggers) == 2


def test_preview_pane_clear_releases_resources(tmp_path: Path, qapp: QApplication):
    """Verify PreviewPane.clear() closes open pymupdf.Document and clears pixmaps and editor buffers."""
    pane = PreviewPane()

    # Load Image
    img_path = tmp_path / "clear_test.png"
    Image.new("RGB", (120, 80), color="red").save(img_path)
    pane.load_result(SearchHit(path=str(img_path), kind=FileKind.IMAGE, score=1.0))
    assert not pane.image_view._pixmap.isNull()

    # Load PDF
    pdf_path = tmp_path / "clear_test.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 50), "Release test page")
    doc.save(str(pdf_path))
    doc.close()
    pane.load_result(SearchHit(path=str(pdf_path), kind=FileKind.PDF, score=1.0, loc=1.0), query="Release")
    pdf_doc = pane.pdf_view._doc
    assert pdf_doc is not None
    assert not pdf_doc.is_closed
    assert not pane.pdf_view._rendered_pixmap.isNull()

    # Load Audio
    pane.load_result(SearchHit(path="/tmp/clear_audio.mp3", kind=FileKind.AUDIO, score=1.0, loc=15.0))
    assert pane.audio_view._path == "/tmp/clear_audio.mp3"

    # Load Video
    pane.load_result(SearchHit(path="/tmp/clear_video.mp4", kind=FileKind.VIDEO, score=1.0, loc=25.0))
    assert not pane.video_view._poster_pixmap.isNull()
    assert pane.video_view._path == "/tmp/clear_video.mp4"

    # Load Code
    py_path = tmp_path / "clear_code.py"
    py_path.write_text("def cleanup():\n    return True\n", encoding="utf-8")
    pane.load_result(SearchHit(path=str(py_path), kind=FileKind.CODE, score=1.0, focus=1.0), query="cleanup")
    assert pane.code_view.editor.toPlainText() != ""
    assert len(pane.code_view.editor.extraSelections()) > 0

    # Clear all views
    pane.clear()

    assert pane.currentIndex() == 0
    assert pane.image_view._pixmap.isNull()
    assert pane.image_view.dimensions_label.text() == ""
    assert pdf_doc.is_closed
    assert pane.pdf_view._doc is None
    assert pane.pdf_view._rendered_pixmap.isNull()
    assert pane.pdf_view._total_pages == 0
    assert pane.pdf_view._current_page == 0
    assert pane.pdf_view._query_terms == []
    assert pane.audio_view._path == ""
    assert pane.audio_view._current_pos == 0.0
    assert pane.video_view._poster_pixmap.isNull()
    assert pane.video_view._path == ""
    assert pane.video_view._current_pos == 0.0
    assert pane.code_view.editor.toPlainText() == ""
    assert pane.code_view.editor.extraSelections() == []



