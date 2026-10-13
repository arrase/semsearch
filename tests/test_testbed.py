"""Unit tests for the synthetic testbed generator."""

import ast
import json
from pathlib import Path
import wave
import docx
from PIL import Image
import pymupdf as fitz
from semsearch.core.search_text import has_words
from semsearch.core.types import SegmentKind
from semsearch.extractors.media import extract_audio, extract_video
from semsearch.extractors.pdfs import extract_pdf
from tests.testbed import (
    create_synthetic_audio,
    create_synthetic_code,
    create_synthetic_digital_pdf,
    create_synthetic_image,
    create_synthetic_markdown_doc,
    create_synthetic_notebook,
    create_synthetic_screenshot,
    create_synthetic_text_doc,
    create_synthetic_video,
    create_synthetic_visual_pdf,
    create_synthetic_word_doc,
    generate_testbed,
)


def test_synthetic_images(tmp_path: Path) -> None:
    jpg_path = create_synthetic_image(tmp_path / "images" / "sunset.jpg")
    assert jpg_path.is_file()
    with Image.open(jpg_path) as img:
        assert img.format == "JPEG"
        assert img.size == (640, 480)
        assert img.mode == "RGB"

    png_path = create_synthetic_screenshot(tmp_path / "images" / "screenshot.png")
    assert png_path.is_file()
    with Image.open(png_path) as img:
        assert img.format == "PNG"
        assert img.size == (1280, 720)
        assert img.mode == "RGBA"


def test_synthetic_pdfs(tmp_path: Path) -> None:
    digital_pdf = create_synthetic_digital_pdf(tmp_path / "pdfs" / "digital.pdf")
    assert digital_pdf.is_file()
    with fitz.open(digital_pdf) as doc:
        assert len(doc) == 1
        text = doc[0].get_text().strip()
        assert len(text) > 200
        assert has_words(text)

    visual_pdf = create_synthetic_visual_pdf(tmp_path / "pdfs" / "visual.pdf")
    assert visual_pdf.is_file()
    with fitz.open(visual_pdf) as doc:
        assert len(doc) == 1
        text = doc[0].get_text().strip()
        assert len(text) < 200

    # Ensure extract_pdf parses digital as text and visual as image without OCR
    segs_digital, comp1 = extract_pdf(str(digital_pdf))
    assert comp1 is True
    assert len(segs_digital) == 1
    assert segs_digital[0].modality.value == "text"

    segs_visual, comp2 = extract_pdf(str(visual_pdf))
    assert comp2 is True
    assert len(segs_visual) == 1
    assert segs_visual[0].modality.value == "image"


def test_synthetic_documents(tmp_path: Path) -> None:
    txt_path = create_synthetic_text_doc(tmp_path / "docs" / "notes.txt")
    assert txt_path.is_file()
    txt_content = txt_path.read_text(encoding="utf-8")
    assert "ERR-4096-FAIL" in txt_content
    assert "INV-2026-X99" in txt_content

    md_path = create_synthetic_markdown_doc(tmp_path / "docs" / "readme.md")
    assert md_path.is_file()
    md_content = md_path.read_text(encoding="utf-8")
    assert md_content.startswith("# SemSearch")
    assert "## Highlights" in md_content

    docx_path = create_synthetic_word_doc(tmp_path / "docs" / "brief.docx")
    assert docx_path.is_file()
    doc = docx.Document(str(docx_path))
    assert len(doc.paragraphs) > 0
    assert "Executive Strategy Brief" in doc.paragraphs[0].text


def test_synthetic_audio_and_video(tmp_path: Path) -> None:
    wav_path = create_synthetic_audio(tmp_path / "media" / "speech.wav", duration_sec=35)
    assert wav_path.is_file()
    with wave.open(str(wav_path), "rb") as wf:
        assert wf.getnchannels() == 1
        assert wf.getframerate() == 16000
        assert wf.getsampwidth() == 2

    audio_windows = extract_audio(str(wav_path))
    assert len(audio_windows) >= 1

    vid_path = create_synthetic_video(tmp_path / "media" / "clip.mp4")
    assert vid_path.is_file()
    segs = extract_video(str(vid_path), sample_interval_sec=3.0)
    frames = [s for s in segs if s.kind == SegmentKind.FRAME]
    audio = [s for s in segs if s.kind == SegmentKind.AUDIO]
    assert len(frames) >= 2
    assert len(audio) >= 1


def test_synthetic_code_and_notebook(tmp_path: Path) -> None:
    py_path = create_synthetic_code(tmp_path / "code" / "service.py")
    assert py_path.is_file()
    py_content = py_path.read_text(encoding="utf-8")
    tree = ast.parse(py_content)
    symbol_names = [node.name for node in tree.body if isinstance(node, (ast.FunctionDef, ast.ClassDef))]
    assert "retry_with_backoff" in symbol_names
    assert "TokenBucket" in symbol_names

    ipynb_path = create_synthetic_notebook(tmp_path / "code" / "analysis.ipynb")
    assert ipynb_path.is_file()
    data = json.loads(ipynb_path.read_text(encoding="utf-8"))
    assert "cells" in data
    code_cells = [c for c in data["cells"] if c.get("cell_type") == "code"]
    assert len(code_cells) > 0
    for cell in code_cells:
        assert cell.get("outputs") == []
        assert "calculate_metrics" in "".join(cell.get("source", []))


def test_generate_testbed_end_to_end(tmp_path: Path) -> None:
    testbed_dir = tmp_path / "testbed"
    files = generate_testbed(testbed_dir)
    assert len(files) == 11
    for key, path in files.items():
        assert path.is_file(), f"File {key} at {path} should exist"
        assert path.stat().st_size > 0
