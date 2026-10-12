"""Unit tests for document extractor and boundary-snapped chunking."""

from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile
from docx import Document
from semsearch.core.types import Modality, SegmentKind
from semsearch.extractors.docs import (
    chunk_document,
    chunk_text,
    extract_doc,
    extract_doc_segments,
    has_words,
    read_doc_text,
    strip_html,
)


def test_has_words() -> None:
    assert has_words("This is standard English text.") is True
    assert has_words("Python 3.13 programming") is True
    assert has_words("1234567890 9876543210") is False
    assert has_words("--- === +++ !!! ???") is False
    assert has_words("   \t\n  ") is False
    # Exactly 20% letters: 2 letters out of 10 non-whitespace chars
    assert has_words("ab12345678") is True
    # 1 letter out of 10 non-whitespace chars: 10%
    assert has_words("a123456789") is False


def test_strip_html() -> None:
    html = "<html><body><h1>Title</h1><p>Paragraph 1 with <b>bold</b> text.</p><p>Paragraph 2.</p></body></html>"
    stripped = strip_html(html)
    assert "Title" in stripped
    assert "Paragraph 1 with bold text." in stripped
    assert "Paragraph 2." in stripped
    assert "<h1>" not in stripped
    assert "<b>" not in stripped


def test_read_doc_text_formats(tmp_path: Path) -> None:
    # 1. Plain text & Markdown
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("Plain text notes line 1\nLine 2", encoding="utf-8")
    assert "Plain text notes" in read_doc_text(str(txt_file))

    md_file = tmp_path / "readme.md"
    md_file.write_text("# Markdown Title\nContent goes here.", encoding="utf-8")
    assert "Markdown Title" in read_doc_text(str(md_file))

    # 2. HTML
    html_file = tmp_path / "doc.html"
    html_file.write_text("<p>Hello from HTML</p>", encoding="utf-8")
    assert "Hello from HTML" in read_doc_text(str(html_file))

    # 3. DOCX
    docx_file = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Docx Heading", level=1)
    doc.add_paragraph("Docx paragraph test content.")
    doc.save(str(docx_file))
    docx_text = read_doc_text(str(docx_file))
    assert "Docx Heading" in docx_text
    assert "Docx paragraph test content." in docx_text

    # 4. ODT
    odt_file = tmp_path / "sample.odt"
    content_xml = b"""<?xml version="1.0" encoding="UTF-8"?>
    <office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
                             xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
      <office:body>
        <office:text>
          <text:h>OpenDocument Title</text:h>
          <text:p>OpenDocument sample body text.</text:p>
        </office:text>
      </office:body>
    </office:document-content>"""
    with zipfile.ZipFile(odt_file, "w") as zf:
        zf.writestr("content.xml", content_xml)
    odt_text = read_doc_text(str(odt_file))
    assert "OpenDocument Title" in odt_text
    assert "OpenDocument sample body text." in odt_text

    # 5. RTF
    rtf_file = tmp_path / "sample.rtf"
    rtf_content = r"{\rtf1\ansi\deff0 {\fonttbl{\f0 Arial;}}\f0\fs24 Hello \b RTF \b0 World!\par Second line.}"
    rtf_file.write_text(rtf_content, encoding="utf-8")
    rtf_text = read_doc_text(str(rtf_file))
    assert "Hello RTF World!" in rtf_text
    assert "Second line." in rtf_text


def test_chunk_document_word_boundaries() -> None:
    # Build text that exceeds 1800 characters
    words = ["word" + str(i) for i in range(500)]
    full_text = " ".join(words)
    assert len(full_text) > 2000

    chunks = chunk_document(full_text, chunk_size=1800, overlap=200)
    assert len(chunks) >= 2
    for chunk in chunks:
        assert len(chunk.text) <= 1800
        # Assert word boundary snapping: no partial words at start or end
        assert not chunk.text.startswith(" ")
        assert not chunk.text.endswith(" ")
        # Every word should be a complete "wordX"
        for w in chunk.text.split():
            assert w.startswith("word")

    # Check chunk_text returns matching tuples
    tuples = chunk_text(full_text, chunk_size=1800, overlap=200)
    assert len(tuples) == len(chunks)
    assert tuples[0][2] == chunks[0].text


def test_extract_doc_segments(tmp_path: Path) -> None:
    doc_path = tmp_path / "test.txt"
    doc_path.write_text("Sentence one about semantic search.\nSentence two about embeddings.", encoding="utf-8")

    segments = extract_doc(str(doc_path))
    assert len(segments) == 1
    seg = segments[0]
    assert seg.kind == SegmentKind.CHUNK
    assert seg.modality == Modality.TEXT
    assert "semantic search" in seg.payload
    assert seg.loc == 0.0


def test_docx_table_extraction(tmp_path: Path) -> None:
    docx_file = tmp_path / "table_sample.docx"
    doc = Document()
    doc.add_heading("Report Summary", level=1)
    doc.add_paragraph("Overview before tables.")

    # Standard table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "98.5%"

    # Table with merged cells (to test clean deduplication)
    table_merged = doc.add_table(rows=2, cols=2)
    table_merged.cell(0, 0).merge(table_merged.cell(0, 1))
    table_merged.cell(0, 0).text = "Merged Header"
    table_merged.cell(1, 0).text = "Column A"
    table_merged.cell(1, 1).text = "Column B"

    doc.save(str(docx_file))

    extracted = read_doc_text(str(docx_file))
    assert "Report Summary" in extracted
    assert "Overview before tables." in extracted
    assert "Metric" in extracted
    assert "Score" in extracted
    assert "Accuracy" in extracted
    assert "98.5%" in extracted
    assert "Merged Header" in extracted
    # Merged header should only appear once despite spanning two columns
    assert extracted.count("Merged Header") == 1
    assert "Column A" in extracted
    assert "Column B" in extracted

    # Also test extract_doc_segments
    segments = extract_doc_segments(str(docx_file))
    assert len(segments) >= 1
    assert any("Merged Header" in s.payload for s in segments)
